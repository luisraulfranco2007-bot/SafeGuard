import os
import openpyxl
from modelo.poliza import Poliza
from modelo.siniestro import Siniestro


COLS_POLIZAS = [
    "numero", "titular", "monto_asegurado",
    "tipo_seguro", "fecha_inicio", "fecha_vencimiento"
]

COLS_SINIESTROS = [
    "id_reporte", "titular", "fecha",
    "descripcion", "numero_poliza"
]


class ManejadorDatos:
    """
    Maneja la lectura y escritura del archivo Excel.
    Es la única clase que toca directamente openpyxl.
    """

    def __init__(self, ruta_archivo: str = "datos/polizas.xlsx"):
        self.__ruta_archivo = ruta_archivo
        self.__asegurar_archivo()

    def __asegurar_archivo(self):
        """Crea el archivo con encabezados si no existe."""
        if not os.path.exists(self.__ruta_archivo):
            os.makedirs(os.path.dirname(self.__ruta_archivo), exist_ok=True)
            wb = openpyxl.Workbook()

            # Hoja Polizas
            ws1 = wb.active
            ws1.title = "Polizas"
            ws1.append(COLS_POLIZAS)

            # Hoja Siniestros
            ws2 = wb.create_sheet("Siniestros")
            ws2.append(COLS_SINIESTROS)

            wb.save(self.__ruta_archivo)

    
    def cargar_polizas(self) -> list:
        """Lee la hoja Polizas y devuelve una lista de objetos Poliza."""
        wb = openpyxl.load_workbook(self.__ruta_archivo)
        ws = wb["Polizas"]
        polizas = []
        for fila in ws.iter_rows(min_row=2, values_only=True):
            if fila[0] is None:
                continue
            numero, titular, monto, tipo, f_ini, f_ven = fila
            polizas.append(
                Poliza(
                    numero=str(numero),
                    titular=str(titular),
                    monto_asegurado=float(monto),
                    tipo_seguro=str(tipo),
                    fecha_inicio=str(f_ini),
                    fecha_vencimiento=str(f_ven),
                )
            )
        return polizas

    def guardar_polizas(self, lista: list):
     
        wb = openpyxl.load_workbook(self.__ruta_archivo)
        ws = wb["Polizas"]

        for fila in ws.iter_rows(min_row=2):
            for celda in fila:
                celda.value = None

        for i, p in enumerate(lista, start=2):
            ws.cell(i, 1, p.get_numero())
            ws.cell(i, 2, p.get_titular())
            ws.cell(i, 3, p.get_monto())
            ws.cell(i, 4, p.get_tipo())
            ws.cell(i, 5, p.get_fecha_inicio())
            ws.cell(i, 6, p.get_fecha_vencimiento())

        wb.save(self.__ruta_archivo)

  
    def cargar_siniestros(self) -> list:
        """Lee la hoja Siniestros y devuelve una lista de objetos Siniestro."""
        wb = openpyxl.load_workbook(self.__ruta_archivo)
        ws = wb["Siniestros"]
        siniestros = []
        for fila in ws.iter_rows(min_row=2, values_only=True):
            if fila[0] is None:
                continue
            id_rep, titular, fecha, desc, num_pol = fila
            siniestros.append(
                Siniestro(
                    id_reporte=str(id_rep),
                    titular=str(titular),
                    fecha=str(fecha),
                    descripcion=str(desc),
                    numero_poliza=str(num_pol),
                )
            )
        return siniestros

    def guardar_siniestros(self, lista: list):
        """Sobreescribe la hoja Siniestros con la lista actual."""
        wb = openpyxl.load_workbook(self.__ruta_archivo)
        ws = wb["Siniestros"]

        for fila in ws.iter_rows(min_row=2):
            for celda in fila:
                celda.value = None

        for i, s in enumerate(lista, start=2):
            ws.cell(i, 1, s.get_id_reporte())
            ws.cell(i, 2, s.get_titular())
            ws.cell(i, 3, s.get_fecha())
            ws.cell(i, 4, s.get_descripcion())
            ws.cell(i, 5, s.get_numero_poliza())

        wb.save(self.__ruta_archivo)

    def siguiente_id_siniestro(self) -> str:
        """Genera el siguiente ID de siniestro en formato SIN-001."""
        siniestros = self.cargar_siniestros()
        return f"SIN-{len(siniestros) + 1:03d}"